from __future__ import annotations

import io
import shutil
import tempfile

from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from docx import Document as DocxDocument

from apps.accounts.models import User, UserRole
from apps.core.models import Institution
from apps.documents.extractors import DocumentTextExtractor, normalize_pdf_pages
from apps.documents.models import Document, DocumentKind

DOCX_MIME = DocumentTextExtractor.DOCX_MIME


def build_docx_bytes(paragraphs: list[str]) -> bytes:
    docx = DocxDocument()

    for paragraph in paragraphs:
        docx.add_paragraph(paragraph)

    buffer = io.BytesIO()
    docx.save(buffer)
    return buffer.getvalue()


class DocxDocumentTestMixin:
    """
    Crea un `Document` DOCX real (archivo en un MEDIA_ROOT temporal) para
    probar la extracción y el pipeline sobre lo que produce python-docx.
    """

    def setUp(self) -> None:
        super().setUp()
        self._media_root = tempfile.mkdtemp()
        self._media_override = override_settings(MEDIA_ROOT=self._media_root)
        self._media_override.enable()

        self.institution = Institution.objects.create(
            name="Instituto de prueba",
            slug=f"instituto-docx-{type(self).__name__.lower()}",
        )
        self.student = User.objects.create_user(
            username=f"alumno-docx-{type(self).__name__.lower()}",
            email=f"alumno-docx-{type(self).__name__.lower()}@example.com",
            password="clave-segura-123",
            role=UserRole.STUDENT,
            institution=self.institution,
        )

    def tearDown(self) -> None:
        self._media_override.disable()
        shutil.rmtree(self._media_root, ignore_errors=True)
        super().tearDown()

    def make_docx_document(self, paragraphs: list[str]) -> Document:
        payload = build_docx_bytes(paragraphs)
        document = Document(
            institution=self.institution,
            owner=self.student,
            uploaded_by=self.student,
            title="Tesis DOCX",
            kind=DocumentKind.OTHER,
            original_filename="tesis.docx",
            mime_type=DOCX_MIME,
            file_size_bytes=len(payload),
            sha256_hash=f"hash-docx-{len(paragraphs)}-{len(payload)}",
        )
        document.original_file.save("tesis.docx", ContentFile(payload), save=False)
        document.save()
        return document


class DocxParagraphSeparationTests(DocxDocumentTestMixin, TestCase):
    """
    Cada párrafo de Word debe quedar separado por "\\n\\n": el filtro, la
    detección de IA y la búsqueda web parten por párrafos. Con "\\n" todo
    el DOCX era un solo bloque (tesis Mantilla: 0 % de similitud web).
    """

    PARAGRAPHS = [
        "Primer párrafo del documento con suficiente texto para contar.",
        "Segundo párrafo del documento, también con texto suficiente.",
        "Tercer párrafo del documento que cierra la prueba de extracción.",
    ]

    def test_docx_paragraphs_are_separated_by_a_blank_line(self) -> None:
        document = self.make_docx_document(self.PARAGRAPHS)

        result = DocumentTextExtractor().extract(document=document)

        self.assertEqual(result.content, "\n\n".join(self.PARAGRAPHS))

    def test_empty_docx_paragraphs_do_not_leave_extra_blank_lines(self) -> None:
        document = self.make_docx_document(
            [self.PARAGRAPHS[0], "", "   ", self.PARAGRAPHS[1], self.PARAGRAPHS[2]],
        )

        result = DocumentTextExtractor().extract(document=document)

        self.assertNotIn("\n\n\n", result.content)
        self.assertEqual(len(result.content.split("\n\n")), 3)


class PdfTextNormalizationTests(TestCase):
    """
    `normalize_pdf_pages` reconstruye párrafos a partir de las líneas
    visuales que devuelve pypdf (una entrada por página).
    """

    HEADER = "Instituto de Educación Superior Pedagógico Público ABL"

    def _page(self, number: int, body: str) -> str:
        return f"{self.HEADER} \n{body}\n{number} "

    def test_lines_of_a_paragraph_are_joined_and_hyphens_removed(self) -> None:
        page = (
            "       El juego simbólico favorece la expresión oral de los niños \n"
            "porque les permite ensayar roles y negociar significados con sus \n"
            "pares en situaciones cotidianas del aula, según la investi-\n"
            "gación revisada en este capítulo. \n"
            "       Además, la docente cumple un rol mediador clave."
        )

        text = normalize_pdf_pages([page])

        self.assertEqual(
            text.split("\n\n"),
            [
                "El juego simbólico favorece la expresión oral de los niños porque les "
                "permite ensayar roles y negociar significados con sus pares en "
                "situaciones cotidianas del aula, según la investigación revisada en "
                "este capítulo.",
                "Además, la docente cumple un rol mediador clave.",
            ],
        )

    def test_page_numbers_and_running_headers_are_removed(self) -> None:
        topics = ["la lectura", "la escritura", "el cálculo", "la oralidad", "el arte", "la ciencia"]
        pages = [
            self._page(
                number,
                f"En esta página se describe {topic} como parte del trabajo docente \n"
                f"cotidiano en el aula, y se explica {topic} con ejemplos concretos \n"
                "de la práctica pedagógica que se observó durante el periodo escolar.",
            )
            for number, topic in zip(range(10, 16), topics)
        ]

        text = normalize_pdf_pages(pages)

        self.assertNotIn(self.HEADER, text)
        self.assertNotRegex(text, r"(?m)^\s*1[0-5]\s*$")
        self.assertIn("trabajo docente cotidiano en el aula", text)

    def test_paragraph_continuing_on_next_page_is_joined(self) -> None:
        pages = [
            self._page(n, body)
            for n, body in enumerate(
                [
                    "Los resultados muestran que la estrategia aplicada mejoró el\n"
                    "nivel de comprensión lectora de los estudiantes del grupo",
                    "experimental frente al grupo de control en todas las dimensiones.",
                    "Texto de relleno de la tercera página con una oración completa.",
                    "Texto de relleno de la cuarta página con una oración completa.",
                ],
                start=20,
            )
        ]

        text = normalize_pdf_pages(pages)

        self.assertIn("estudiantes del grupo experimental frente al grupo", text)

    def test_heading_leads_its_section_and_short_lines_stay_together(self) -> None:
        page = (
            "Dedicatoria \n"
            "A mis padres, por su apoyo incondicional durante todos estos años de \n"
            "estudio y por enseñarme el valor del esfuerzo y la constancia diaria. \n"
            "\n"
            "Índice \n"
            "1. Introducción ........................ 5 \n"
            "2. Marco teórico ....................... 9 \n"
            "3. Metodología ......................... 21 "
        )

        blocks = normalize_pdf_pages([page]).split("\n\n")

        self.assertTrue(blocks[0].startswith("Dedicatoria\nA mis padres, por su apoyo"))
        self.assertEqual(blocks[1].split("\n")[0], "Índice")
        self.assertEqual(len(blocks[1].split("\n")), 4)

    def test_numbered_items_start_new_blocks(self) -> None:
        page = (
            "Los objetivos específicos de la investigación son los siguientes y se \n"
            "detallan a continuación en el orden en que fueron evaluados. \n"
            "1. Identificar el nivel de expresión oral antes de la intervención. \n"
            "2. Identificar el nivel de expresión oral después de la intervención."
        )

        text = normalize_pdf_pages([page])

        self.assertIn("\n1. Identificar", text.replace("\n\n", "\n"))
        self.assertIn("\n2. Identificar", text)
